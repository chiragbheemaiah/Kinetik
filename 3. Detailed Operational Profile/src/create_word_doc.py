import json
import argparse
import csv
from pathlib import Path
from typing import Any, Optional, List

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

# -------------------------
# DEFAULT CONFIG PATHS
# -------------------------
INPUT_JSON_DEFAULT = Path("outputs/profile_detailed.json")
OUTPUT_DOCX_DEFAULT = Path("outputs/enterprise_operational_profile.docx")

TOP_LEVEL_HEADING_STYLE = "Heading 1"
SECOND_LEVEL_HEADING_STYLE = "Heading 2"
THIRD_LEVEL_HEADING_STYLE = "Heading 3"
BODY_STYLE = "Normal"
TABLE_STYLE = "Table Grid"

# -------------------------
# STYLE HELPERS
# -------------------------

def ensure_custom_styles(doc: Document) -> None:
    styles = doc.styles
    normal_style = styles[BODY_STYLE]
    if normal_style.type == WD_STYLE_TYPE.PARAGRAPH:
        font = normal_style.font
        font.name = "Avenir Light"
        font.size = Pt(10)
        font.color.rgb = RGBColor(0, 0, 0)

# -------------------------
# RENDERING HELPERS
# -------------------------

def add_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    if not text:
        return
    p = doc.add_paragraph(text)
    if style:
        p.style = style


def _sanitize_filename_part(s: str) -> str:
    cleaned = []
    for ch in s.lower():
        if ch.isalnum():
            cleaned.append(ch)
        else:
            cleaned.append("_")
    out = "".join(cleaned)
    while "__" in out:
        out = out.replace("__", "_")
    return out.strip("_") or "section"


def export_table_csv(items: List[dict], csv_dir: Path, trail: List[str]) -> Optional[Path]:
    try:
        csv_dir.mkdir(parents=True, exist_ok=True)
        # Build filename from trail
        name = "_".join(_sanitize_filename_part(t) for t in trail)
        path = csv_dir / f"{name}.csv"
        # Collect headers as union preserving order by first row, then others
        headers = list(items[0].keys()) if items else []
        seen = set(headers)
        for row in items[1:]:
            for k in row.keys():
                if k not in seen:
                    headers.append(k)
                    seen.add(k)
        with path.open("w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=headers)
            writer.writeheader()
            for row in items:
                writer.writerow({k: row.get(k, "") for k in headers})
        return path
    except Exception:
        return None


def render_value(doc: Document, key: str, value: Any, level: int = 1, trail: Optional[List[str]] = None, csv_dir: Optional[Path] = None) -> None:
    if value is None:
        return

    heading_style = None
    if isinstance(value, (str, list, dict)):
        if level == 1:
            heading_style = TOP_LEVEL_HEADING_STYLE
        elif level == 2:
            heading_style = SECOND_LEVEL_HEADING_STYLE
        elif level == 3:
            heading_style = THIRD_LEVEL_HEADING_STYLE

        if heading_style:
            add_paragraph(doc, key, style=heading_style)
        else:
            if not (isinstance(value, str) and value.strip() == ""):
                p = doc.add_paragraph()
                run = p.add_run(f"{key}: ")
                run.bold = True

    if isinstance(value, str):
        text = value.strip()
        if text:
            add_paragraph(doc, text, style=BODY_STYLE)
    elif isinstance(value, list):
        render_list(doc, value, level + 1)
    elif isinstance(value, dict):
        render_dict(doc, value, level + 1, (trail or []) + [key], csv_dir)


def render_list(doc: Document, items: list, level: int) -> None:
    if not items:
        return

    if all(isinstance(i, (str, int, float, bool)) or i is None for i in items):
        for item in items:
            if item is None:
                continue
            p = doc.add_paragraph(str(item))
            p.style = BODY_STYLE
            p.style = "List Bullet"
        return

    dict_items = [obj for obj in items if isinstance(obj, dict)]
    if dict_items and len(dict_items) == len(items):
        headers = list(dict_items[0].keys())
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = TABLE_STYLE
        hdr_cells = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = str(h)
        for row in dict_items:
            cells = table.add_row().cells
            for i, h in enumerate(headers):
                val = row.get(h, "")
                cells[i].text = "" if val is None else str(val)
        doc.add_paragraph("")
        return

    for obj in items:
        if isinstance(obj, dict):
            doc.add_paragraph("")
            for k, v in obj.items():
                if v is None:
                    continue
                if isinstance(v, (str, int, float, bool)):
                    p = doc.add_paragraph()
                    run = p.add_run(f"{k}: ")
                    run.bold = True
                    p.add_run(str(v))
                else:
                    render_value(doc, k, v, level)
        else:
            p = doc.add_paragraph(str(obj))
            p.style = BODY_STYLE


def render_dict(doc: Document, d: dict, level: int, trail: Optional[List[str]] = None, csv_dir: Optional[Path] = None) -> None:
    for k, v in d.items():
        if isinstance(v, str) and v.strip() == "":
            continue
        # CSV export if this is a Table with list-of-dicts
        if csv_dir is not None and isinstance(v, list) and (k.lower() == "table") and v and all(isinstance(x, dict) for x in v):
            export_table_csv(v, csv_dir, (trail or []) + [k])
        render_value(doc, k, v, level, trail, csv_dir)


def insert_table_of_contents(doc: Document) -> None:
    add_paragraph(doc, "Table of Contents", style=TOP_LEVEL_HEADING_STYLE)
    p = doc.add_paragraph()
    # Build a complex field for TOC: TOC \o "1-3" \h \z \u
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')

    instr_text = OxmlElement('w:instrText')
    instr_text.set(qn('xml:space'), 'preserve')
    instr_text.text = 'TOC \\o "1-3" \\h \\z \\u'

    fld_separate = OxmlElement('w:fldChar')
    fld_separate.set(qn('w:fldCharType'), 'separate')

    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')

    r = OxmlElement('w:r')
    r.append(fld_begin)
    r.append(instr_text)
    r.append(fld_separate)
    r2 = OxmlElement('w:r')
    r2.append(fld_end)
    p._p.append(r)
    p._p.append(r2)

# -------------------------
# MAIN
# -------------------------

def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate Word document from JSON profile")
    parser.add_argument("--input", "-i", type=str, default=str(INPUT_JSON_DEFAULT), help="Path to input JSON file")
    parser.add_argument("--output", "-o", type=str, default=str(OUTPUT_DOCX_DEFAULT), help="Path to output DOCX file")
    parser.add_argument("--export-csv-dir", type=str, default=None, help="Optional directory to export any embedded tables as CSVs")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    input_path = Path(args.input)
    output_path = Path(args.output)
    csv_dir = Path(args.export_csv_dir) if args.export_csv_dir else None

    if not input_path.exists():
        raise FileNotFoundError(f"Input JSON not found: {input_path}")

    with input_path.open("r", encoding="utf-8") as f:
        profile = json.load(f)

    doc = Document()
    ensure_custom_styles(doc)

    doc.add_heading("Citigroup Inc. – Operational and Technology Profile", level=0)
    doc.add_paragraph("Generated from JSON profile.", style=BODY_STYLE)
    insert_table_of_contents(doc)
    doc.add_page_break()

    for top_key, top_val in profile.items():
        render_value(doc, top_key, top_val, level=1, trail=[top_key], csv_dir=csv_dir)
        doc.add_paragraph("")

    output_path.parent.mkdir(exist_ok=True)
    doc.save(output_path)
    print(f"Saved Word document to: {output_path}")


if __name__ == "__main__":
    main()
